"""
Web version of the face attendance system.
Put this file in the SAME folder as attendance_cv.py, then:
    pip install flask
    python app.py        ->  open http://localhost:5000
Login password: set ADMIN_PASSWORD env var (default: admin123 - change it!)
"""
import base64
import calendar
import json
import os
import secrets
import sys
import time
import urllib.request
from getpass import getpass
from datetime import date, datetime
from pathlib import Path
from threading import Lock

import cv2
import numpy as np
from flask import (Flask, abort, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash

from attendance_cv import (EXCEL_PATH, best_match, db, get_feature,
                           load_known, load_models, log_to_excel, mark)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret")
PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123")
detector, recognizer = load_models(640, 480)
lock = Lock()  # the models are not thread-safe
Path("captures").mkdir(exist_ok=True)
REG_DIR = Path(os.environ.get("DATA_DIR", ".")) / "registers"
REG_DIR.mkdir(exist_ok=True)


def reg_path(year, month):
    return REG_DIR / f"register_{year}-{month:02d}.xlsx"


def cell_str(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_students(file):
    """Read (name, student_id) pairs from an uploaded sheet with a header row."""
    ws = load_workbook(file, data_only=True).active
    rows = [r for r in ws.iter_rows(values_only=True) if r and any(r)]
    head = [str(c or "").lower() for c in rows[0]]
    ni = next((i for i, h in enumerate(head) if "name" in h), 0)
    ii = next((i for i, h in enumerate(head) if "id" in h or "index" in h), 1)
    return [(cell_str(r[ni]), cell_str(r[ii])) for r in rows[1:]
            if len(r) > max(ni, ii) and r[ni] and r[ii]]


def add_row(ws, name, sid, ndays):
    r = ws.max_row + 1
    ws.cell(r, 1, name)
    ws.cell(r, 2, sid)
    ws.cell(r, 3 + ndays, f'=COUNTIF(C{r}:{get_column_letter(2 + ndays)}{r},"P")')


def create_register(year, month, students):
    ndays = calendar.monthrange(year, month)[1]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{calendar.month_name[month]} {year}"
    ws.append(["Name", "Staff ID"] + [date(year, month, d) for d in range(1, ndays + 1)]
              + ["Total"])
    for c in range(3, 3 + ndays):
        ws.cell(1, c).number_format = "ddd d-mmm"
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 14
    for name, sid in students:
        add_row(ws, name, sid, ndays)
    wb.save(reg_path(year, month))


def fill_absent(ws, upto):
    """Write A in empty cells for every day before day number `upto`."""
    for r in range(2, ws.max_row + 1):
        for d in range(1, upto):
            if ws.cell(r, 2 + d).value is None:
                ws.cell(r, 2 + d, "A")


def mark_register(name, sid, now):
    path = reg_path(now.year, now.month)
    if not path.exists():
        return "missing"
    try:
        wb = load_workbook(path)
        ws = wb.active
        ndays = calendar.monthrange(now.year, now.month)[1]
        row = next((r for r in range(2, ws.max_row + 1)
                    if cell_str(ws.cell(r, 2).value) == sid), None)
        if row is None:
            add_row(ws, name, sid, ndays)
            row = ws.max_row
        ws.cell(row, 2 + now.day, "P")
        fill_absent(ws, now.day)
        wb.save(path)
        return "ok"
    except PermissionError:
        return "locked"


fails = {}  # ip -> (failed attempts, last time)


def locked():
    n, t = fails.get(request.remote_addr, (0, 0))
    return n >= 5 and time.time() - t < 300


def add_fail():
    n, _ = fails.get(request.remote_addr, (0, 0))
    fails[request.remote_addr] = (n + 1, time.time())


def admin_row():
    con = db()
    con.execute("CREATE TABLE IF NOT EXISTS admin (id INTEGER PRIMARY KEY CHECK (id = 1), "
                "pw_hash TEXT, rc_hash TEXT)")
    con.commit()
    return con.execute("SELECT pw_hash, rc_hash FROM admin WHERE id = 1").fetchone()


def check_pw(pw):
    row = admin_row()
    return check_password_hash(row[0], pw) if row else pw == PASSWORD


ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "").strip()
RESEND_KEY = os.environ.get("RESEND_API_KEY", "").strip()
MAIL_FROM = os.environ.get("MAIL_FROM", "FaceAttend <onboarding@resend.dev>")


def save_password(pw):
    con = db()
    con.execute("INSERT OR REPLACE INTO admin VALUES (1, ?, NULL)", (generate_password_hash(pw),))
    con.commit()


def mask_email(e):
    name, _, host = e.partition("@")
    return name[:1] + "***@" + host


def send_mail(subject, html, verbose=False):
    """Email ADMIN_EMAIL through Resend's HTTPS API. Returns an error text, or None if sent."""
    if not (ADMIN_EMAIL and RESEND_KEY):
        return "Email is not set up (ADMIN_EMAIL or RESEND_API_KEY is missing)."
    req = urllib.request.Request(
        "https://api.resend.com/emails",
        data=json.dumps({"from": MAIL_FROM, "to": [ADMIN_EMAIL],
                         "subject": subject, "html": html}).encode(),
        headers={"Authorization": f"Bearer {RESEND_KEY}",
                 "Content-Type": "application/json", "User-Agent": "FaceAttend/1.0"})
    try:
        urllib.request.urlopen(req, timeout=15)
    except Exception as e:
        detail = e.read().decode()[:300] if hasattr(e, "read") else str(e)
        print("Email error:", detail)
        return f"Could not send the email. {detail}" if verbose else "Could not send the email."


def reset_con():
    con = db()
    con.execute("CREATE TABLE IF NOT EXISTS reset (id INTEGER PRIMARY KEY CHECK (id = 1), "
                "code_hash TEXT, expires REAL, tries INTEGER)")
    return con


def start_reset():
    """Email a 6-digit code valid for 15 minutes. Returns an error text or None."""
    con = reset_con()
    row = con.execute("SELECT expires FROM reset").fetchone()
    if row and row[0] - 900 > time.time() - 60:
        return "A code was sent a moment ago. Wait a minute, then try again."
    code = f"{secrets.randbelow(10 ** 6):06d}"
    err = send_mail("Your FaceAttend reset code",
                    f"<p>Your password reset code is:</p><h2>{code}</h2>"
                    "<p>It expires in 15 minutes. If you didn't ask for it, ignore this email.</p>")
    if err:
        return err
    con.execute("INSERT OR REPLACE INTO reset VALUES (1, ?, ?, 0)",
                (generate_password_hash(code), time.time() + 900))
    con.commit()


def finish_reset(code, new, confirm):
    """Check the emailed code and set the new password. Returns an error text or None."""
    con = reset_con()
    row = con.execute("SELECT code_hash, expires, tries FROM reset").fetchone()
    if not row or time.time() > row[1] or row[2] >= 5:
        return "That code has expired. Send a new one."
    if not check_password_hash(row[0], code.strip()):
        con.execute("UPDATE reset SET tries = tries + 1")
        con.commit()
        return "Wrong code."
    err = pw_error(new, confirm)
    if not err:
        save_password(new)
        con.execute("DELETE FROM reset")
        con.commit()
    return err


def pw_error(new, confirm):
    if len(new) < 8:
        return "Use at least 8 characters."
    if new != confirm:
        return "The two passwords don't match."


def page(template, **ctx):
    return render_template(template, **ctx)


def decode(data_url):
    raw = base64.b64decode(data_url.split(",")[1])
    img = cv2.imdecode(np.frombuffer(raw, np.uint8), cv2.IMREAD_COLOR)
    h, w = img.shape[:2]
    if w > 640:
        img = cv2.resize(img, (640, int(h * 640 / w)))
    return img


def detect(img):
    h, w = img.shape[:2]
    detector.setInputSize((w, h))
    return detector.detect(img)[1]


@app.before_request
def guard():
    if request.endpoint not in ("login", "forgot", "static") and not session.get("ok"):
        return redirect("/login")


@app.route("/login", methods=["GET", "POST"])
def login():
    err = ""
    if request.method == "POST":
        if locked():
            err = "Too many attempts. Try again in 5 minutes."
        elif check_pw(request.form.get("pw", "")):
            fails.pop(request.remote_addr, None)
            session["ok"] = True
            return redirect("/")
        else:
            add_fail()
            err = "Wrong password."
    return page("login.html", err=err)


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.route("/account", methods=["GET", "POST"])
def account():
    msg = err = None
    if request.method == "POST":
        f = request.form
        if not check_pw(f.get("current", "")):
            err = "Current password is wrong."
        elif f.get("action") == "test":
            err = send_mail("FaceAttend test email", "<p>Email is working.</p>", verbose=True)
            msg = None if err else "Test email sent to " + mask_email(ADMIN_EMAIL) + "."
        else:
            err = pw_error(f.get("new", ""), f.get("confirm", ""))
            if not err:
                save_password(f["new"])
                msg = "Password changed."
    return page("account.html", msg=msg, err=err,
                email=mask_email(ADMIN_EMAIL) if ADMIN_EMAIL else "")


@app.route("/forgot", methods=["GET", "POST"])
def forgot():
    if session.get("ok"):
        return redirect("/account")
    msg = err = None
    step = "send"
    if request.method == "POST":
        f = request.form
        if locked():
            err = "Too many attempts. Try again in 5 minutes."
        elif f.get("action") == "send":
            err = start_reset()
            step = "send" if err else "code"
            msg = None if err else "Code sent to " + mask_email(ADMIN_EMAIL) + ". It expires in 15 minutes."
        else:
            step = "code"
            err = finish_reset(f.get("code", ""), f.get("new", ""), f.get("confirm", ""))
            if err == "Wrong code.":
                add_fail()
            if not err:
                fails.pop(request.remote_addr, None)
                step, msg = "done", "Password reset. You can sign in now."
    return page("forgot.html", msg=msg, err=err, step=step,
                email=mask_email(ADMIN_EMAIL) if ADMIN_EMAIL else "")


@app.route("/")
def dashboard():
    date = request.args.get("date") or datetime.now().strftime("%Y-%m-%d")
    con = db()
    rows = con.execute(
        """SELECT s.name, a.student_id, a.time FROM attendance a
           JOIN students s ON s.student_id = a.student_id
           WHERE a.date = ? ORDER BY a.time""", (date,)).fetchall()
    studs = con.execute(
        """SELECT s.name, s.student_id, COUNT(p.id) FROM students s
           LEFT JOIN samples p ON p.student_id = s.student_id
           GROUP BY s.student_id ORDER BY s.name""").fetchall()
    now = datetime.now()
    cur = reg_path(now.year, now.month)
    regs = [p.stem.replace("register_", "")
            for p in sorted(REG_DIR.glob("register_*.xlsx"), reverse=True)]
    return page("dashboard.html", date=date, rows=rows, studs=studs, msg=request.args.get("msg"),
                ym=now.strftime("%Y-%m"), month_label=now.strftime("%B %Y"),
                active=cur.name if cur.exists() else None, regs=regs)


@app.route("/attend")
def attend():
    return page("attend.html")


@app.route("/enroll")
def enroll_page():
    return page("enroll.html")


@app.post("/delete/<sid>")
def delete(sid):
    con = db()
    for table in ("samples", "attendance", "students"):
        con.execute(f"DELETE FROM {table} WHERE student_id = ?", (sid,))
    con.commit()
    return redirect("/")


@app.route("/export")
def export():
    if not Path(EXCEL_PATH).exists():
        abort(404, "No attendance.xlsx yet - mark someone present first.")
    return send_file(EXCEL_PATH, as_attachment=True)


@app.post("/register/upload")
def register_upload():
    try:
        year, month = map(int, request.form["month"].split("-"))
        students = read_students(request.files["file"])
    except Exception:
        return redirect(url_for("dashboard", msg="Could not read that file. Use an .xlsx "
                                "with Name and Staff ID columns."))
    path = reg_path(year, month)
    if not students:
        msg = "No staff found in that sheet."
    elif path.exists() and not request.form.get("replace"):
        msg = f"A register for {year}-{month:02d} already exists. Tick 'replace' to overwrite it."
    else:
        create_register(year, month, students)
        msg = f"Register for {year}-{month:02d} created for {len(students)} staff."
    return redirect(url_for("dashboard", msg=msg))


@app.route("/register/<ym>")
def register_download(ym):
    year, month = map(int, ym.split("-"))
    path = reg_path(year, month)
    if not path.exists():
        abort(404)
    now = datetime.now()
    if (year, month) == (now.year, now.month):
        upto = now.day
    elif (year, month) < (now.year, now.month):
        upto = calendar.monthrange(year, month)[1] + 1
    else:
        upto = 1
    try:
        wb = load_workbook(path)
        fill_absent(wb.active, upto)
        wb.save(path)
    except PermissionError:
        pass  # file open in Excel: download as is
    return send_file(path, as_attachment=True)


@app.post("/api/enroll")
def api_enroll():
    d = request.json
    name, sid = d.get("name", "").strip(), d.get("sid", "").strip()
    if not name or not sid:
        return jsonify(ok=False, msg="Name and staff ID are required.")
    img = decode(d["image"])
    with lock:
        faces = detect(img)
        if faces is None or len(faces) != 1:
            return jsonify(ok=False, msg="Need exactly one face in view. Try again.")
        feat = get_feature(recognizer, img, faces[0])
    con = db()
    if d["n"] == 0:  # first capture replaces any older enrollment
        con.execute("DELETE FROM samples WHERE student_id = ?", (sid,))
        con.execute("INSERT OR REPLACE INTO students VALUES (?,?)", (sid, name))
    con.execute("INSERT INTO samples (student_id, feature) VALUES (?,?)",
                (sid, feat.astype(np.float32).tobytes()))
    con.commit()
    return jsonify(ok=True, msg=f"Captured sample {d['n'] + 1}/3")


@app.post("/api/recognize")
def api_recognize():
    img = decode(request.json["image"])
    con = db()
    with lock:
        faces = detect(img)
        if faces is None:
            return jsonify(status="none", msg="No face detected")
        face = max(faces, key=lambda r: r[2] * r[3])  # largest face
        feat = get_feature(recognizer, img, face)
        names, ids, feats = load_known(con)
        sid, score = best_match(recognizer, feat, ids, feats)
    if not sid:
        return jsonify(status="unknown", msg=f"Unknown face ({score:.2f})")
    if mark(con, sid):
        now = datetime.now()
        log_to_excel(names[sid], sid, now)
        reg = mark_register(names[sid], sid, now)
        cv2.imwrite(f"captures/{now:%Y-%m-%d_%H%M%S}_{sid}.jpg", img)
        note = {"missing": " (no register for this month - upload one)",
                "locked": " (close the register in Excel)"}.get(reg, "")
        return jsonify(status="present", msg=f"{names[sid]} - PRESENT{note}")
    return jsonify(status="already", msg=f"{names[sid]} - already marked today")


if __name__ == "__main__":
    if sys.argv[1:] == ["reset-password"]:  # emergency reset from this computer
        new = getpass("New password: ")
        error = pw_error(new, getpass("Confirm: "))
        if not error:
            save_password(new)
        print(error or "Password reset. You can sign in now.")
    else:
        app.run(host="127.0.0.1", port=5000, threaded=True)
